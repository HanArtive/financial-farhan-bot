import os
from dotenv import load_dotenv
from supabase import create_client
from openpyxl import load_workbook

load_dotenv()

# =========================
# SUPABASE
# =========================

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

supabase = create_client(
    SUPABASE_URL,
    SUPABASE_KEY
)

# =========================
# EXCEL ONEDRIVE
# =========================

FILE_EXCEL = r"C:\Users\FARHAN\OneDrive\Financial_Plan_Farhan_DENSO_6_Bulan.xlsx"

SHEET_NAME = "Transaksi"
TABLE_NAME = "TransaksiTable"


def sync_excel():

    print("===================================")
    print("   SUPABASE → EXCEL SYNC")
    print("===================================")

    # =========================
    # AMBIL DATA SUPABASE
    # =========================

    response = (
        supabase
        .table("transaksi")
        .select("*")
        .order("id")
        .execute()
    )

    supabase_data = response.data

    print(f"Data Supabase : {len(supabase_data)} transaksi")

    # =========================
    # BUKA EXCEL
    # =========================

    if not os.path.exists(FILE_EXCEL):
        raise FileNotFoundError(
            f"File Excel tidak ditemukan:\n{FILE_EXCEL}"
        )

    workbook = load_workbook(FILE_EXCEL)

    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise Exception(
            f"Sheet '{SHEET_NAME}' tidak ditemukan."
        )

    sheet = workbook[SHEET_NAME]

    # =========================
    # BACA DATA EXCEL BERDASARKAN ID
    # =========================

    excel_rows = {}

    for row in range(2, sheet.max_row + 1):

        transaction_id = sheet.cell(row, 1).value

        if transaction_id is not None:
            excel_rows[str(transaction_id)] = row

    # =========================
    # ID YANG ADA DI SUPABASE
    # =========================

    supabase_ids = set()

    # =========================
    # INSERT / UPDATE
    # =========================

    for transaction in supabase_data:

        transaction_id = str(transaction["id"])

        supabase_ids.add(transaction_id)

        if transaction_id in excel_rows:

            # =====================
            # UPDATE BARIS LAMA
            # =====================

            row = excel_rows[transaction_id]

            sheet.cell(row, 2).value = transaction.get("tanggal")
            sheet.cell(row, 3).value = transaction.get("jenis")
            sheet.cell(row, 4).value = transaction.get("kategori")
            sheet.cell(row, 5).value = transaction.get("nominal")
            sheet.cell(row, 6).value = transaction.get("keterangan")

            print(
                f"UPDATE → ID {transaction_id}"
            )

        else:

            # =====================
            # TAMBAH BARIS BARU
            # =====================

            row = sheet.max_row + 1

            sheet.cell(row, 1).value = transaction["id"]
            sheet.cell(row, 2).value = transaction.get("tanggal")
            sheet.cell(row, 3).value = transaction.get("jenis")
            sheet.cell(row, 4).value = transaction.get("kategori")
            sheet.cell(row, 5).value = transaction.get("nominal")
            sheet.cell(row, 6).value = transaction.get("keterangan")

            excel_rows[transaction_id] = row

            print(
                f"INSERT → ID {transaction_id}"
            )

        # Format tanggal
        sheet.cell(row, 2).number_format = "dd/mm/yyyy hh:mm"

        # Format nominal
        sheet.cell(row, 5).number_format = '#,##0'

    # =========================
    # DELETE
    # =========================

    rows_to_delete = []

    for transaction_id, row in excel_rows.items():

        if transaction_id not in supabase_ids:

            rows_to_delete.append(row)

    # Hapus dari bawah supaya nomor baris tidak bergeser
    for row in sorted(rows_to_delete, reverse=True):

        transaction_id = sheet.cell(row, 1).value

        print(
            f"DELETE → ID {transaction_id}"
        )

        sheet.delete_rows(row, 1)

    # =========================
    # UPDATE TABLE RANGE
    # =========================

    if TABLE_NAME in sheet.tables:

        table = sheet.tables[TABLE_NAME]

        last_row = max(1, sheet.max_row)

        table.ref = f"A1:F{last_row}"

    # =========================
    # SIMPAN
    # =========================

    workbook.save(FILE_EXCEL)

    workbook.close()

    print("-----------------------------------")
    print("✅ Sinkronisasi selesai!")
    print(f"📊 Supabase : {len(supabase_data)} transaksi")
    print(f"📁 Excel    : {FILE_EXCEL}")
    print("-----------------------------------")


if __name__ == "__main__":
    sync_excel()