from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path

FILE_EXCEL = Path(__file__).parent / "Financial_Plan_Farhan_DENSO_6_Bulan.xlsx"


def add_transaction(jenis, kategori, nominal, keterangan=""):
    workbook = load_workbook(FILE_EXCEL)

    if "Transaksi" not in workbook.sheetnames:
        worksheet = workbook.create_sheet("Transaksi")

        worksheet.append([
            "ID",
            "Tanggal",
            "Jenis",
            "Kategori",
            "Nominal",
            "Keterangan"
        ])
    else:
        worksheet = workbook["Transaksi"]

    # Menentukan ID transaksi berikutnya
    if worksheet.max_row <= 1:
        transaction_id = 1
    else:
        transaction_id = worksheet.max_row

    tanggal = datetime.now()

    worksheet.append([
        transaction_id,
        tanggal,
        jenis,
        kategori,
        nominal,
        keterangan
    ])

    workbook.save(FILE_EXCEL)

    return transaction_id, tanggal