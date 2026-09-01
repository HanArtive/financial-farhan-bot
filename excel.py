import os
from datetime import datetime

from dotenv import load_dotenv
from supabase import create_client
from openpyxl import load_workbook

load_dotenv()

# =========================
# SUPABASE
# =========================

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

supabase = create_client(
    SUPABASE_URL,
    SUPABASE_KEY
)

# =========================
# EXCEL ONEDRIVE
# =========================

FILE_EXCEL = r"C:\Users\FARHAN\OneDrive\Financial_Plan_Farhan_DENSO_6_Bulan.xlsx"

SHEET_NAME = "Transaksi"
TABLE_NAME = "TransaksiTable"


def add_transaction(jenis, kategori, nominal, keterangan=""):

    # =========================
    # 1. SIMPAN KE SUPABASE
    # =========================

    response = (
        supabase
        .table("transaksi")
        .insert({
            "jenis": jenis,
            "kategori": kategori,
            "nominal": nominal,
            "keterangan": keterangan
        })
        .execute()
    )

    if not response.data:
        raise Exception("Transaksi gagal disimpan ke Supabase.")

    transaction = response.data[0]

    transaction_id = transaction["id"]

    # =========================
    # 2. SIMPAN KE EXCEL
    # =========================

    tanggal = datetime.now()

    if not os.path.exists(FILE_EXCEL):
        raise FileNotFoundError(
            f"File Excel tidak ditemukan:\n{FILE_EXCEL}"
        )

    workbook = load_workbook(FILE_EXCEL)

    if SHEET_NAME not in workbook.sheetnames:
        raise Exception(
            f"Sheet '{SHEET_NAME}' tidak ditemukan."
        )

    sheet = workbook[SHEET_NAME]

    # Cari baris kosong berikutnya
    next_row = sheet.max_row + 1

    # Isi data sesuai kolom Transaksi
    sheet.cell(next_row, 1).value = transaction_id
    sheet.cell(next_row, 2).value = tanggal
    sheet.cell(next_row, 3).value = jenis
    sheet.cell(next_row, 4).value = kategori
    sheet.cell(next_row, 5).value = nominal
    sheet.cell(next_row, 6).value = keterangan

    # Format tanggal
    sheet.cell(next_row, 2).number_format = "dd/mm/yyyy hh:mm"

    # Format nominal Rupiah
    sheet.cell(next_row, 5).number_format = '#,##0'

    # =========================
    # 3. PERPANJANG EXCEL TABLE
    # =========================

    if TABLE_NAME in sheet.tables:
        table = sheet.tables[TABLE_NAME]

        table.ref = (
            f"A1:F{next_row}"
        )

    # Simpan kembali ke Excel OneDrive
    workbook.save(FILE_EXCEL)

    workbook.close()

    return transaction_id, tanggal